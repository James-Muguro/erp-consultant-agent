"""
Extracts plain text from uploaded project documents so their content can
be fed into the per-project memory store (project_memories, see
src/memory/project_memory.py) and become part of what agents recall.

Deliberately simple: no OCR, no layout/table reconstruction, no chunking
strategy beyond a flat character cap. Matches the "keyword/tag search, no
embeddings" scope decided for the memory layer - a full semantic ingestion
pipeline isn't in scope here.
"""
from io import BytesIO
from typing import Optional

from pypdf import PdfReader
from docx import Document as DocxDocument

from src.utils.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Caps how much extracted text lands in one project_memories row. Not a
# real chunking strategy (a very long document is simply truncated, not
# split into multiple searchable entries) - good enough for the
# keyword-search scope decided for this stage; revisit if long documents
# turn out to need their later sections to actually be findable.
MAX_EXTRACTED_CHARS = 20_000


class UnsupportedFileType(Exception):
    pass


def extract_text(filename: str, content: bytes) -> str:
    """Returns extracted text (possibly empty, e.g. for a scanned/image-only
    PDF with no text layer - that's a valid, expected outcome, not an
    error). Raises UnsupportedFileType for anything outside
    SUPPORTED_EXTENSIONS - callers turn that into a 415."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    try:
        if ext == ".pdf":
            text = _extract_pdf(content)
        elif ext == ".docx":
            text = _extract_docx(content)
        else:  # .txt, .md
            text = content.decode("utf-8", errors="replace")
    except Exception as e:
        # A corrupt/malformed file shouldn't crash the upload - it should
        # come back as "uploaded, but nothing extracted", same as a
        # scanned PDF with no text layer. The file itself is still saved.
        logger.warning("Text extraction failed, continuing with empty text", filename=filename, error=str(e))
        return ""

    return text[:MAX_EXTRACTED_CHARS]


def _extract_pdf(content: bytes) -> str:
    reader = PdfReader(BytesIO(content))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages.append(page_text)
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    doc = DocxDocument(BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                paragraphs.append(row_text)
    return "\n".join(paragraphs)
