"""
TOR extraction tests.

Snapshot basis: src/services/tor_extraction.py as seen during Phase
2.5/2.6 inspection.

The extractor is deterministic. It parses:
  * Word (.docx) via python-docx: table rows preferred, paragraph
    fallback otherwise.
  * Excel (.xlsx) via openpyxl: heuristic longest-cell selection.

Coverage:
  * Word table extraction produces one row per data row.
  * Word paragraph fallback when the doc has no tables.
  * Excel extraction produces one row per data row.
  * Unknown format raises TorExtractionError.
  * Sequential external_code generation (TOR-001, TOR-002, ...).
  * Header rows are skipped.
  * detect_format returns the expected canonical value.
"""
from __future__ import annotations

import io

import pytest

from src.services import tor_extraction


def _make_word_table(rows):
    """Build a .docx in memory with a single table."""
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=0, cols=3)
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_word_paragraphs(texts):
    from docx import Document
    doc = Document()
    for t in texts:
        doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_excel(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestDetectFormat:
    @pytest.mark.parametrize("filename,ct,expected", [
        ("tender.docx", "", "word"),
        ("tender.doc", "", "word"),
        ("tender.xlsx", "", "excel"),
        ("tender.xls", "", "excel"),
        ("unknown", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "word"),
        ("unknown", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "excel"),
        ("unknown.txt", "text/plain", None),
    ])
    def test_detects_format_from_filename_or_content_type(
        self, filename, ct, expected,
    ):
        assert tor_extraction.detect_format(filename, ct) == expected


class TestWordExtraction:
    def test_table_rows_are_extracted(self):
        content = _make_word_table([
            ["ID", "Description", "Category"],
            ["1", "Bank reconciliation required", "Finance"],
            ["2", "Purchase approval workflow", "Procurement"],
        ])
        rows = tor_extraction.extract_tor(content, "word")
        assert len(rows) == 2
        assert rows[0]["external_code"] == "TOR-001"
        assert rows[1]["external_code"] == "TOR-002"
        assert "Bank reconciliation" in rows[0]["description"]
        assert rows[0]["category"] == "Finance"

    def test_paragraph_fallback_when_no_tables(self):
        content = _make_word_paragraphs([
            "This is a long paragraph describing the tender background "
            "and context for the implementation project.",
            "The solution must support multi-currency transactions "
            "across every legal entity in scope.",
            "abc",  # too short, skipped
        ])
        rows = tor_extraction.extract_tor(content, "word")
        assert len(rows) == 2
        assert "multi-currency" in rows[1]["description"]


class TestExcelExtraction:
    def test_excel_rows_are_extracted(self):
        content = _make_excel([
            ["ID", "Description", "Category"],
            [1, "Approval workflow for purchase requisitions", "Procurement"],
            [2, "Automated period close checklist", "Finance"],
        ])
        rows = tor_extraction.extract_tor(content, "excel")
        assert len(rows) == 2
        assert rows[0]["external_code"] == "TOR-001"
        assert "Approval workflow" in rows[0]["description"]


class TestErrorHandling:
    def test_unknown_format_raises(self):
        with pytest.raises(tor_extraction.TorExtractionError):
            tor_extraction.extract_tor(b"junk", "csv")

    def test_corrupt_word_raises(self):
        with pytest.raises(tor_extraction.TorExtractionError):
            tor_extraction.extract_tor(b"not a docx", "word")

    def test_corrupt_excel_raises(self):
        with pytest.raises(tor_extraction.TorExtractionError):
            tor_extraction.extract_tor(b"not an xlsx", "excel")